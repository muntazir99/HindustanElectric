"""GST summary for the CA (plan §16.3). An 18% item at ₹118 (₹100 + ₹18 GST) keeps every figure checkable by hand."""

import io
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from catalog.models import Item, Product
from sales.models import Customer

pytestmark = pytest.mark.django_db
D = Decimal


def priced(make_item, product, price, gst, hsn=""):
    item = make_item(product=product)
    Product.objects.filter(pk=item.product_id).update(gst_rate=D(gst), hsn_code=hsn)
    Item.objects.filter(pk=item.pk).update(selling_price=D(price))
    return item


def sell(client, item, quantity, **bill):
    unit = item.units.get(is_base=True)
    draft = client.post("/api/sales/invoices", {"lines": [{"item": item.pk, "unit": unit.pk, "quantity": str(quantity)}], **bill},
                        format="json").data
    return client.post(f"/api/sales/invoices/{draft['id']}/finalise", {"payments": [{"mode": "cash", "amount": draft["total"]}]},
                       format="json").data


def give_back(client, bill, quantity):
    client.post(f"/api/sales/invoices/{bill['id']}/returns",
                {"lines": [{"line": bill["lines"][0]["id"], "quantity": str(quantity)}], "refund_mode": "cash", "reason": "Not needed"},
                format="json")


@pytest.fixture
def month(owner_api, owner, make_item):
    switch = priced(make_item, "Switch", "118", "18", hsn="8536")
    bulb = priced(make_item, "Free bulb", "50", "0")  # 0%, and no HSN entered
    trader = Customer.objects.create(name="Sharma Traders", gstin="10ABCDE1234F1Z5", created_by=owner)
    b2b = sell(owner_api, switch, 2, customer=trader.pk)  # ₹236
    walk_in = sell(owner_api, switch, 3)  # ₹354
    sell(owner_api, switch, 1000, place_of_supply="09", buyer_name="UP Contractor", buyer_address="Gorakhpur")  # ₹1,18,000
    cancelled = sell(owner_api, switch, 1)
    owner_api.post(f"/api/sales/invoices/{cancelled['id']}/cancel", {"reason": "Wrong item"}, format="json")
    sell(owner_api, bulb, 1)
    give_back(owner_api, b2b, 1)
    give_back(owner_api, walk_in, 1)
    estimate = owner_api.post("/api/sales/invoices", {"lines": [{"item": switch.pk, "unit": switch.units.get().pk, "quantity": "5"}]},
                              format="json").data
    owner_api.post(f"/api/sales/invoices/{estimate['id']}/quotation", format="json")
    return b2b


def this_month(client, extra=""):
    today = timezone.localdate()
    return client.get(f"/api/sales/gst?from={today.replace(day=1)}&to={today}{extra}")


def test_tax_by_rate_is_net_of_returns(owner_api, month):
    summary = this_month(owner_api).json()
    [nil, eighteen] = summary["by_rate"]
    # 18%: 200 (GSTIN) + 300 (walk-in) + 1,00,000 (Uttar Pradesh) − 100 − 100 returned.
    assert eighteen == {"rate": "18", "taxable": "100300.00", "igst": "18000.00", "cgst": "27.00", "sgst": "27.00", "gst": "18054.00"}
    assert nil == {"rate": "0", "taxable": "50.00", "igst": "0.00", "cgst": "0.00", "sgst": "0.00", "gst": "0.00"}
    assert summary["counts"] == {"bills": 4, "b2b": 1, "b2cl": 1, "returns": 2, "cdnr": 1}  # the cancelled bill and estimate are out
    assert summary["missing_hsn"] == ["Havells Free bulb"]
    [bills, notes] = summary["documents"]
    assert (bills["Total Number"], bills["Cancelled"], notes["Total Number"]) == (5, 1, 2)
    assert bills["Sr. No. From"].endswith("/00001") and bills["Sr. No. To"].endswith("/00005")


def test_excel_in_gstr1_layout(owner_api, month):
    book = load_workbook(io.BytesIO(this_month(owner_api, "&download=xlsx").content))
    assert book.sheetnames == ["Summary", "b2b", "b2cl", "b2cs", "cdnr", "exemp", "hsn(b2b)", "hsn(b2c)", "docs"]

    def rows(name):
        header, *data = book[name].iter_rows(values_only=True)
        return [dict(zip(header, row)) for row in data]

    [b2b] = rows("b2b")
    assert (b2b["GSTIN/UIN of Recipient"], b2b["Invoice Number"], b2b["Invoice Value"], b2b["Place Of Supply"]) == (
        "10ABCDE1234F1Z5", month["number"], 236, "10-Bihar")
    assert (b2b["Rate"], b2b["Taxable Value"], b2b["Reverse Charge"], b2b["Invoice Type"]) == (18, 200, "N", "Regular B2B")
    [b2cl] = rows("b2cl")
    assert (b2cl["Place Of Supply"], b2cl["Invoice Value"], b2cl["Taxable Value"]) == ("09-Uttar Pradesh", 118000, 100000)
    [b2cs] = rows("b2cs")  # walk-in 300 − 100 returned; the big UP bill is in b2cl, not here
    assert (b2cs["Type"], b2cs["Place Of Supply"], b2cs["Rate"], b2cs["Taxable Value"]) == ("OE", "10-Bihar", 18, 200)
    [cdnr] = rows("cdnr")
    assert (cdnr["Note Type"], cdnr["Note Value"], cdnr["Taxable Value"], cdnr["GSTIN/UIN of Recipient"]) == ("C", 118, 100, "10ABCDE1234F1Z5")
    [nil] = rows("exemp")
    assert (nil["Description"], nil["Nil Rated Supplies"]) == ("Intra-State supplies to unregistered persons", 50)

    [hsn_b2b] = rows("hsn(b2b)")
    assert (hsn_b2b["HSN"], hsn_b2b["UQC"], hsn_b2b["Total Quantity"], hsn_b2b["Taxable Value"], hsn_b2b["Total Value"]) == (
        "8536", "PCS-PIECES", 1, 100, 118)
    switch = next(row for row in rows("hsn(b2c)") if row["HSN"] == "8536")
    assert (switch["Total Quantity"], switch["Taxable Value"], switch["Integrated Tax Amount"], switch["Central Tax Amount"]) == (
        1002, 100200, 18000, 18)


def test_owner_only_and_a_quarter_at_most(staff_api, owner_api):
    assert staff_api.get("/api/sales/gst").status_code == 403
    assert owner_api.get("/api/sales/gst?from=2026-04-01&to=2026-09-30").status_code == 400


def test_last_month_by_default(owner_api):
    first = timezone.localdate().replace(day=1)
    summary = owner_api.get("/api/sales/gst").json()
    assert summary["to"] < first.isoformat() and summary["from"].endswith("-01")
