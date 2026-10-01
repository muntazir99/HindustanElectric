import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from catalog.importers import CATALOGUE_COLUMNS
from catalog.models import Brand, Item, ItemUnit, Product

pytestmark = pytest.mark.django_db
D = Decimal

HEADERS = [header for header, _, _ in CATALOGUE_COLUMNS]


def xlsx(rows, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(
        "items.xlsx", output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def row(**values):
    keys = [key for _, key, _ in CATALOGUE_COLUMNS]
    return [values.get(key, "") for key in keys]


RED = row(category="Wires & Cables", brand="Havells", product="Lifeline wire", variant="1.5 sq mm Red",
          hsn="8544", gst=18, unit="mtr", selling_price=28, pack_unit="coil", pack_size=90,
          pack_barcode=8901234500011, pack_price=2400, rack="W1", min_stock=180)
SWITCH = row(category="Switches", brand="Anchor", product="Roma switch", variant="6A White", hsn=8536,
             unit="pcs", mrp=45, selling_price=38, barcode="8901234500028", rack="S2", aliases="button")


def upload(client, file, kind="catalogue", **flags):
    return client.post(f"/api/import/{kind}", {"file": file, **flags}, format="multipart")


class TestCatalogueImport:
    def test_staff_cannot_import(self, staff_api):
        assert upload(staff_api, xlsx([RED])).status_code == 403
        assert staff_api.get("/api/import/catalogue/template").status_code == 403

    def test_template_has_the_columns(self, owner_api):
        response = owner_api.get("/api/import/catalogue/template")
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        assert [cell.value for cell in sheet[1]] == HEADERS

    def test_preview_changes_nothing(self, owner_api):
        response = upload(owner_api, xlsx([RED, SWITCH]))
        assert response.status_code == 200
        assert response.data["summary"] == {
            "rows": 2, "created": 2, "updated": 0, "unchanged": 0, "errors": 0, "committed": False,
        }
        assert not Item.objects.exists() and not Brand.objects.exists()
        assert "code" not in response.data["rows"][0]  # real codes are given only when saved

    def test_commit_creates_items_units_and_barcodes(self, owner_api):
        response = upload(owner_api, xlsx([RED, SWITCH]), commit="true")
        assert response.data["summary"]["committed"] is True
        red = Item.objects.get(variant="1.5 sq mm Red")
        assert (red.base_unit, red.selling_price, red.rack, red.min_stock) == ("m", D("28"), "W1", D("180"))
        coil = red.units.get(name="coil")
        assert (coil.factor, coil.barcode, coil.selling_price) == (D("90"), "8901234500011", D("2400"))
        switch = Item.objects.get(variant="6A White")
        assert switch.product.hsn_code == "8536"
        assert switch.units.get(is_base=True).barcode == "8901234500028"
        assert "button" in switch.search_text

    def test_reimport_updates_instead_of_duplicating(self, owner_api):
        upload(owner_api, xlsx([RED]), commit="true")
        changed = list(RED)
        changed[HEADERS.index("Selling price")] = 29
        response = upload(owner_api, xlsx([changed]), commit="true")
        assert response.data["summary"]["updated"] == 1
        assert Item.objects.count() == 1
        assert Item.objects.get().selling_price == D("29")

    def test_blank_cells_leave_values_alone(self, owner_api):
        upload(owner_api, xlsx([RED]), commit="true")
        upload(owner_api, xlsx([row(brand="Havells", product="Lifeline wire", variant="1.5 sq mm Red", rack="W9")]), commit="true")
        item = Item.objects.get()
        assert (item.rack, item.selling_price, item.base_unit) == ("W9", D("28"), "m")

    def test_row_errors_are_reported_and_block_commit(self, owner_api):
        rows = [
            RED,
            row(brand="Havells", product=""),
            row(product="Fan", unit="dozen"),
            row(product="Bulb", gst=15),
            row(product="Tape", barcode="8901234500011"),  # same as RED's pack barcode
            RED,  # duplicate row
        ]
        response = upload(owner_api, xlsx(rows), commit="true")
        summary = response.data["summary"]
        assert (summary["errors"], summary["committed"]) == (5, False)
        messages = {r["row"]: " ".join(r.get("messages", [])) for r in response.data["rows"]}
        assert "Product name is missing" in messages[3]
        assert "Unknown unit" in messages[4]
        assert "GST" in messages[5]
        assert "already used by" in messages[6]
        assert "Same product and variant as row 2" in messages[7]
        assert not Item.objects.exists()

    def test_skip_errors_saves_good_rows(self, owner_api):
        response = upload(owner_api, xlsx([RED, row(product="Fan", unit="dozen")]), commit="true", skip_errors="true")
        assert response.data["summary"]["committed"] is True
        assert Item.objects.count() == 1

    def test_csv_upload(self, owner_api):
        text = ",".join(HEADERS) + "\nSwitches,Anchor,Roma switch,6A White,8536,18,pc,45,38,,,,,,,S2,,\n"
        file = SimpleUploadedFile("items.csv", text.encode(), content_type="text/csv")
        response = upload(owner_api, file, commit="true")
        assert response.data["summary"]["created"] == 1

    def test_unknown_headers_give_clear_error(self, owner_api):
        response = upload(owner_api, xlsx([["a", "b"]], headers=["Foo", "Bar"]))
        assert response.status_code == 400
        assert "No known column headings" in response.data["detail"]

    def test_existing_product_gets_new_variant(self, owner_api):
        upload(owner_api, xlsx([RED]), commit="true")
        black = list(RED)
        black[HEADERS.index("Variant")] = "1.5 sq mm Black"
        black[HEADERS.index("Pack barcode")] = "8901234500099"
        upload(owner_api, xlsx([black]), commit="true")
        assert Product.objects.count() == 1
        assert Item.objects.count() == 2


class TestPriceImport:
    def price_file(self, rows):
        return xlsx(rows, headers=["Barcode or code", "MRP", "Selling price"])

    def test_updates_base_and_pack_prices_with_preview(self, owner_api):
        upload(owner_api, xlsx([RED, SWITCH]), commit="true")
        switch = Item.objects.get(variant="6A White")
        file = self.price_file([[switch.code, 48, 40], [8901234500011, "", 2500], ["nope", 1, 1]])
        preview = upload(owner_api, file, kind="prices")
        assert preview.data["summary"]["updated"] == 2
        assert preview.data["summary"]["errors"] == 1
        assert "Selling price 38.00 → 40" in preview.data["rows"][0]["messages"]
        switch.refresh_from_db()
        assert switch.selling_price == D("38")  # preview only

        file = self.price_file([[switch.code, 48, 40], [8901234500011, "", 2500]])
        upload(owner_api, file, kind="prices", commit="true")
        switch.refresh_from_db()
        assert (switch.mrp, switch.selling_price) == (D("48"), D("40"))
        assert ItemUnit.objects.get(barcode="8901234500011").selling_price == D("2500")

    def test_unknown_import_kind(self, owner_api):
        assert owner_api.get("/api/import/nonsense/template").status_code == 404
