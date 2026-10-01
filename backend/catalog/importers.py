"""
Excel / CSV import for the catalogue and for price updates.

Every import runs in one transaction with a savepoint per row, so a preview
(commit=False) shows exactly what would happen, including clashes with
existing data, and then rolls everything back.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import GST_RATES, Item, ItemUnit, Product, hsn_validator
from .services import (
    CatalogError,
    brand_named,
    category_named,
    check_barcode_free,
    create_item,
    find_product,
    parse_base_unit,
    set_base_unit,
    set_pack,
)

MAX_ROWS = 10000

# (header, key, help text)
CATALOGUE_COLUMNS = [
    ("Category", "category", "e.g. Wires & Cables. Created if new."),
    ("Brand", "brand", "e.g. Havells. Created if new. Blank for unbranded goods."),
    ("Product", "product", "Required. The product line, e.g. Lifeline Plus HRFR wire."),
    ("Variant", "variant", "Size / colour / rating, e.g. 1.5 sq mm Red. Blank if none."),
    ("HSN", "hsn", "4, 6 or 8 digit HSN code."),
    ("GST %", "gst", "0, 5, 12, 18, 28 or 40. Default 18."),
    ("Unit", "unit", "Smallest unit you sell in: pc, m, ft, kg, l, set, pair, pkt, roll. Default pc."),
    ("MRP", "mrp", "MRP per unit."),
    ("Selling price", "selling_price", "Your price per unit, GST included."),
    ("Barcode", "barcode", "Barcode of a single unit, if it has one."),
    ("Pack unit", "pack_unit", "If also bought/sold in packs: coil, box, bundle..."),
    ("Pack size", "pack_size", "Units in one pack, e.g. 90 (m in a coil) or 100 (pc in a box)."),
    ("Pack barcode", "pack_barcode", "Barcode printed on the pack."),
    ("Pack MRP", "pack_mrp", "MRP of the whole pack."),
    ("Pack price", "pack_price", "Your price for the whole pack, GST included."),
    ("Rack", "rack", "Where it is kept, e.g. A3."),
    ("Min stock", "min_stock", "Reorder when stock falls to this (in units)."),
    ("Aliases", "aliases", "Other names customers use, comma separated."),
]

PRICE_COLUMNS = [
    ("Barcode or code", "code", "Item barcode, pack barcode, or the item code shown in the app."),
    ("MRP", "mrp", "New MRP for that unit or pack. Blank = unchanged."),
    ("Selling price", "selling_price", "New selling price for that unit or pack, GST included. Blank = unchanged."),
]

CATALOGUE_EXAMPLES = [
    ["Wires & Cables", "Havells", "Lifeline Plus HRFR wire", "1.5 sq mm Red", "8544", "18", "m",
     "", "28", "", "coil", "90", "8901234500011", "2650", "2400", "W1", "180", ""],
    ["Switches", "Anchor", "Roma 6A switch", "White", "8536", "18", "pc",
     "45", "38", "8901234500028", "box", "20", "8901234500035", "", "", "S2", "40", "button switch"],
]

TEXT_KEYS = {"barcode", "pack_barcode", "code", "hsn"}


# --- reading files -----------------------------------------------------------

def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_rows(upload, columns):
    """Return [(row_number, {key: text})] from an .xlsx or .csv upload."""
    name = (upload.name or "").lower()
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        raw = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise CatalogError(f"Could not read the Excel file: {exc}") from exc
        raw = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise CatalogError("Upload an .xlsx (Excel) or .csv file.")

    header_index = next((i for i, row in enumerate(raw) if any(_cell_text(c) for c in row)), None)
    if header_index is None:
        raise CatalogError("The file is empty.")
    by_header = {header.lower(): key for header, key, _ in columns}
    keys = [by_header.get(_cell_text(cell).lower()) for cell in raw[header_index]]
    if not any(keys):
        expected = ", ".join(header for header, _, _ in columns)
        raise CatalogError(f"No known column headings found. Expected: {expected}.")

    rows = []
    for offset, row in enumerate(raw[header_index + 1:], start=header_index + 2):
        values = {key: "" for _, key, _ in columns}
        for key, cell in zip(keys, row):
            if key:
                values[key] = _cell_text(cell)
        if any(values.values()):
            rows.append((offset, values))
    if len(rows) > MAX_ROWS:
        raise CatalogError(f"Too many rows ({len(rows)}). Split the file into parts of {MAX_ROWS}.")
    return rows


def _decimal(row, key, label):
    text = row.get(key, "").replace(",", "").replace("₹", "").strip()
    if not text:
        return None
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise CatalogError(f'{label}: "{row[key]}" is not a number.') from exc
    if value < 0:
        raise CatalogError(f"{label} can't be negative.")
    return value


def _gst(row):
    value = _decimal(row, "gst", "GST %")
    if value is None:
        return None
    if value not in {rate for rate, _ in GST_RATES}:
        raise CatalogError(f"GST % must be 0, 5, 12, 18, 28 or 40 (got {row['gst']}).")
    return value


def _messages(exc):
    if isinstance(exc, DjangoValidationError):
        if hasattr(exc, "message_dict"):
            return [f"{field}: {msg}" for field, msgs in exc.message_dict.items() for msg in msgs]
        return list(exc.messages)
    if isinstance(exc, IntegrityError):
        return ["Conflicts with existing data (duplicate name or barcode)."]
    return [str(exc)]


def _run(rows, handle_row, *, commit, skip_errors):
    results = []
    with transaction.atomic():
        for number, row in rows:
            try:
                with transaction.atomic():
                    results.append({"row": number, **handle_row(number, row)})
            except (CatalogError, DjangoValidationError, IntegrityError) as exc:
                label = " ".join(filter(None, [row.get("brand"), row.get("product"), row.get("variant"), row.get("code")]))
                results.append({"row": number, "status": "error", "item": label, "messages": _messages(exc)})
        errors = sum(1 for result in results if result["status"] == "error")
        committed = commit and (errors == 0 or skip_errors)
        if not committed:
            transaction.set_rollback(True)
            # New items get their real code only when actually saved.
            for result in results:
                if result["status"] == "created":
                    result.pop("code", None)
    summary = {
        "rows": len(results),
        "created": sum(1 for r in results if r["status"] == "created"),
        "updated": sum(1 for r in results if r["status"] == "updated"),
        "unchanged": sum(1 for r in results if r["status"] == "unchanged"),
        "errors": errors,
        "committed": committed,
    }
    return {"summary": summary, "rows": results}


# --- catalogue -----------------------------------------------------------------

def import_catalogue(upload, *, commit=False, skip_errors=False):
    rows = read_rows(upload, CATALOGUE_COLUMNS)
    seen = {}

    def handle(number, row):
        if not row["product"]:
            raise CatalogError("Product name is missing.")
        key = (row["brand"].lower(), row["product"].lower(), row["variant"].lower())
        if key in seen:
            raise CatalogError(f"Same product and variant as row {seen[key]}.")
        seen[key] = number
        return _catalogue_row(row)

    return _run(rows, handle, commit=commit, skip_errors=skip_errors)


def _catalogue_row(row):
    base_unit = parse_base_unit(row["unit"]) if row["unit"] else None
    gst = _gst(row)
    hsn = row["hsn"].replace(" ", "")
    if hsn:
        hsn_validator(hsn)
    brand = brand_named(row["brand"])
    category = category_named(row["category"])

    product = find_product(row["product"], brand)
    if product is None:
        product = Product(name=row["product"].strip(), brand=brand, category=category, hsn_code=hsn)
        if gst is not None:
            product.gst_rate = gst
        product.full_clean()
        product.save()
    else:
        changed = False
        for field, value in (("category", category), ("hsn_code", hsn or None), ("gst_rate", gst)):
            if value is not None and getattr(product, field) != value:
                setattr(product, field, value)
                changed = True
        if changed:
            product.full_clean()
            product.save()

    fields = {}
    for key, label in (("mrp", "MRP"), ("selling_price", "Selling price"), ("min_stock", "Min stock")):
        value = _decimal(row, key, label)
        if value is not None:
            fields[key] = value
    for key in ("rack", "aliases"):
        if row[key]:
            fields[key] = row[key]

    pack = None
    if row["pack_unit"] or row["pack_size"]:
        size = _decimal(row, "pack_size", "Pack size")
        if not row["pack_unit"] or not size:
            raise CatalogError("Give both Pack unit and Pack size, or neither.")
        pack = {
            "name": row["pack_unit"],
            "factor": size,
            "barcode": row["pack_barcode"] or None,
            "mrp": _decimal(row, "pack_mrp", "Pack MRP"),
            "selling_price": _decimal(row, "pack_price", "Pack price"),
        }

    item = product.items.filter(variant__iexact=row["variant"].strip()).first()
    if item is None:
        item = create_item(
            product, row["variant"], base_unit or "pc", barcode=row["barcode"] or None, pack=pack, **fields
        )
        return {"status": "created", "item": item.name, "code": item.code}

    if base_unit:
        set_base_unit(item, base_unit)
    for field, value in fields.items():
        setattr(item, field, value)
    item.full_clean(exclude=["product"])
    item.save()
    if row["barcode"]:
        base = item.units.get(is_base=True)
        base.barcode = check_barcode_free(row["barcode"], base)
        base.save()
    if pack:
        set_pack(item, pack["name"], pack["factor"], barcode=pack["barcode"], mrp=pack["mrp"],
                 selling_price=pack["selling_price"])
    return {"status": "updated", "item": item.name, "code": item.code}


# --- prices ---------------------------------------------------------------------

def import_prices(upload, *, commit=False, skip_errors=False):
    rows = read_rows(upload, PRICE_COLUMNS)
    return _run(rows, _price_row, commit=commit, skip_errors=skip_errors)


def _price_row(number, row):
    code = row["code"]
    if not code:
        raise CatalogError("Barcode or code is missing.")
    unit = ItemUnit.objects.select_related("item__product__brand").filter(barcode=code).first()
    if unit is None:
        item = Item.objects.select_related("product__brand").filter(code__iexact=code).first()
        if item is None:
            raise CatalogError(f"No item with barcode or code {code}.")
        unit = item.units.get(is_base=True)
    item = unit.item
    target = item if unit.is_base else unit
    changes = []
    for key, label in (("mrp", "MRP"), ("selling_price", "Selling price")):
        value = _decimal(row, key, label)
        if value is None:
            continue
        old = getattr(target, key)
        if old != value:
            changes.append(f"{label} {old if old is not None else '—'} → {value}")
            setattr(target, key, value)
    label = item.name if unit.is_base else f"{item.name} ({unit.name})"
    if not changes:
        return {"status": "unchanged", "item": label, "code": item.code}
    target.save()
    return {"status": "updated", "item": label, "code": item.code, "messages": changes}


# --- templates ------------------------------------------------------------------

def build_template(columns, examples=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Items"
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    for index, (header, key, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = max(12, len(header) + 4)
        if key in TEXT_KEYS:
            # Long barcodes lose digits if Excel treats them as numbers.
            for row in range(2, 2001):
                sheet.cell(row=row, column=index).number_format = "@"
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("How to fill")
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 80
    guide.append(["Column", "What to enter"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for header, _, help_text in columns:
        guide.append([header, help_text])
    guide.append([])
    guide.append(["Notes", "Fill the Items sheet only. One row per item (each size/colour is its own row)."])
    guide.append(["", "Blank cells leave existing values unchanged when the item already exists."])
    if examples:
        guide.append([])
        guide.append(["Examples"])
        guide.append([header for header, _, _ in columns])
        for example in examples:
            guide.append(example)
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
